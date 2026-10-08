import io

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.agregacao import Agregador, numero
from app.main import criar_app

SPEC = {
    "swagger": "2.0",
    "info": {"title": "teste", "version": "0"},
    "paths": {
        "/sia": {
            "get": {
                "parameters": [
                    {"name": "co_ibge", "in": "query", "type": "integer"},
                    {"name": "limit", "in": "query", "type": "integer", "description": "Deve ser menor ou igual 4."},
                    {"name": "offset", "in": "query", "type": "integer", "default": 0},
                ]
            }
        }
    },
}
# 10 linhas de produção; valores em formatos diferentes, como aparecem em APIs reais
PRODUCAO = [
    {"ds_procedimento": "CONSULTA", "ds_complex": "AB", "qt_procedimento": 10, "nu_valor_procedimento": "100.50"},
    {"ds_procedimento": "CONSULTA", "ds_complex": "MC", "qt_procedimento": "5", "nu_valor_procedimento": 49.5},
    {"ds_procedimento": "RAIO X", "ds_complex": "MC", "qt_procedimento": 2, "nu_valor_procedimento": "1.000,00"},
    {"ds_procedimento": "CONSULTA", "ds_complex": "AB", "qt_procedimento": 1, "nu_valor_procedimento": "10,25"},
    {"ds_procedimento": "EXAME", "ds_complex": "MC", "qt_procedimento": 7, "nu_valor_procedimento": None},
    {"ds_procedimento": "EXAME", "ds_complex": "MC", "qt_procedimento": "n/d", "nu_valor_procedimento": 3},
    {"ds_procedimento": "RAIO X", "ds_complex": "MC", "qt_procedimento": 3, "nu_valor_procedimento": 1500},
    {"ds_procedimento": None, "ds_complex": "AB", "qt_procedimento": 1, "nu_valor_procedimento": 1},
    {"ds_procedimento": "CONSULTA", "ds_complex": "AB", "qt_procedimento": 4, "nu_valor_procedimento": 40},
    {"ds_procedimento": "EXAME", "ds_complex": "AC", "qt_procedimento": 1, "nu_valor_procedimento": 0.5},
]


@pytest.fixture
def api(config):
    def handler(req):
        if req.url.path == "/spec.json":
            return httpx.Response(200, json=SPEC)
        lim, off = int(req.url.params["limit"]), int(req.url.params["offset"])
        return httpx.Response(200, json={"sia": PRODUCAO[off * lim:(off + 1) * lim]})

    with TestClient(criar_app(config, transport=httpx.MockTransport(handler))) as c:
        yield c


def test_numero():
    assert numero(10) == 10 and numero("5") == 5 and numero("100.50") == 100.5
    assert numero("1.000,00") == 1000 and numero("10,25") == 10.25 and numero("1.234.567") == 1234567
    assert numero(None) is None and numero("") is None and numero("n/d") is None and numero(True) is None


def test_agregador_direto():
    ag = Agregador(["a"], ["v"])
    for r in [{"a": "x", "v": 1}, {"a": "y", "v": 5}, {"a": "x", "v": "2"}]:
        ag.adicionar(r)
    res = ag.resultado()
    assert res["linhas"] == [{"a": "y", "registros": 1, "soma_v": 5}, {"a": "x", "registros": 2, "soma_v": 3}]
    assert res["totais"] == {"registros": 3, "soma_v": 8}


def test_soma_de_producao_por_procedimento(api):
    r = api.get(
        "/api/datasets/sia/dados",
        params={"agrupar_por": "ds_procedimento", "somar": "qt_procedimento,nu_valor_procedimento", "max_registros": 0},
    ).json()
    assert r["paginas_consultadas"] == 4 and r["registros_buscados"] == 10  # 4+4+2 e a vazia
    ag = r["agregado"]
    assert ag["colunas"] == ["ds_procedimento", "registros", "soma_qt_procedimento", "soma_nu_valor_procedimento"]
    por_proc = {linha["ds_procedimento"]: linha for linha in ag["linhas"]}
    assert por_proc["CONSULTA"] == {
        "ds_procedimento": "CONSULTA", "registros": 4, "soma_qt_procedimento": 20, "soma_nu_valor_procedimento": 200.25
    }
    assert por_proc["RAIO X"]["soma_nu_valor_procedimento"] == 2500
    assert por_proc[""]["registros"] == 1  # categoria vazia vira ""
    assert ag["linhas"][0]["ds_procedimento"] == "CONSULTA"  # ordenado pela 1ª soma, decrescente
    assert ag["totais"] == {"registros": 10, "soma_qt_procedimento": 34, "soma_nu_valor_procedimento": 2704.75}
    assert any("1 valor(es) não numérico(s) em 'qt_procedimento'" in a for a in r["avisos"])


def test_categorias_sem_soma_e_soma_sem_categoria(api):
    r = api.get("/api/datasets/sia/dados", params={"agrupar_por": "ds_complex,ds_procedimento", "max_registros": 0}).json()
    linhas = r["agregado"]["linhas"]
    assert {"ds_complex": "AB", "ds_procedimento": "CONSULTA", "registros": 3} in linhas
    r = api.get("/api/datasets/sia/dados", params={"somar": "qt_procedimento", "max_registros": 0}).json()
    assert r["agregado"]["linhas"] == [{"registros": 10, "soma_qt_procedimento": 34}]


def test_variavel_inexistente_avisa(api):
    r = api.get("/api/datasets/sia/dados", params={"agrupar_por": "uf", "max_registros": 0}).json()
    assert any("'uf' não existe" in a for a in r["avisos"])


def test_csv_e_xlsx_agregados(api):
    params = {"agrupar_por": "ds_procedimento", "somar": "qt_procedimento", "max_registros": 0}
    r = api.get("/api/datasets/sia/dados", params={**params, "formato": "csv", "separador": ";"})
    linhas = r.content.decode("utf-8-sig").strip().split("\n")
    assert linhas[0] == "ds_procedimento;registros;soma_qt_procedimento"
    assert linhas[-1] == "TOTAL;10;34"
    r = api.get("/api/datasets/sia/dados", params={**params, "formato": "xlsx"})
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["resumo", "dados", "consulta"]
    resumo = list(wb["resumo"].values)
    assert resumo[0] == ("ds_procedimento", "registros", "soma_qt_procedimento")
    assert resumo[-1] == ("TOTAL", 10, 34)
    assert len(list(wb["dados"].values)) == 11


def test_detalhe_omitido_acima_do_limite(config, api):
    api.app.state.config.max_linhas_detalhe = 5
    r = api.get("/api/datasets/sia/dados", params={"somar": "qt_procedimento", "max_registros": 0}).json()
    assert r["dados"] == [] and r["agregado"]["totais"]["registros"] == 10
    assert any("detalhados foram omitidos" in a for a in r["avisos"])


def test_cli_resumo_com_agregado(tmp_path, monkeypatch, config):
    from app import cli
    from app.cliente import ClienteDataSUS

    def handler(req):
        if req.url.path == "/spec.json":
            return httpx.Response(200, json=SPEC)
        lim, off = int(req.url.params["limit"]), int(req.url.params["offset"])
        return httpx.Response(200, json={"sia": PRODUCAO[off * lim:(off + 1) * lim]})

    original = ClienteDataSUS.__init__
    monkeypatch.setattr(ClienteDataSUS, "__init__", lambda self, cfg, transport=None: original(self, cfg, httpx.MockTransport(handler)))
    monkeypatch.setattr(cli.Config, "do_ambiente", classmethod(lambda cls: config))
    saida, resumo = tmp_path / "p.csv", tmp_path / "r.md"
    assert cli.main(["dados", "sia", "-g", "ds_procedimento", "-s", "nu_valor_procedimento", "-n", "0",
                     "-o", str(saida), "--separador", ";", "--resumo", str(resumo)]) == 0
    assert saida.read_text(encoding="utf-8-sig").splitlines()[-1] == "TOTAL;10;2704.75"
    assert (tmp_path / "p_dados.csv").exists()
    texto = resumo.read_text()
    assert "### Resumo por `ds_procedimento`" in texto
    assert "| RAIO X | 2 | 2.500 |" in texto and "| **TOTAL** | **10** | **2.704,75** |" in texto
