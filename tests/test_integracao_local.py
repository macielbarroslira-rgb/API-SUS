"""Fluxo completo do app local pelo criar_app real (sem navegador):

baixar uma base da API oficial (simulada) -> listar -> colunas -> categorias -> consulta com filtros
e agregação -> exportar (Excel e CSV) -> excluir. Os números são conferidos com um cálculo independente.
"""

import asyncio
import csv
import io
import json
import re
import time
from collections import Counter, defaultdict
from urllib.parse import unquote

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import criar_app

FINAIS = {"concluido", "erro", "cancelado"}

SPEC = {
    "swagger": "2.0",
    "info": {"title": "Especificação SINTÉTICA para o teste de integração", "version": "0.0-teste"},
    "basePath": "/",
    "paths": {
        "/sia/producao": {
            "get": {
                "tags": ["Assistência"],
                "summary": "Produção ambulatorial (sintética)",
                "parameters": [
                    {"name": "ano", "in": "query", "type": "integer", "required": True},
                    {"name": "limit", "in": "query", "type": "integer", "maximum": 20},
                    {"name": "offset", "in": "query", "type": "integer"},
                ],
                "responses": {"200": {"description": "ok"}},
            }
        }
    },
}

VALORES = ["1.234,56", "10,5", "100.25", "7", None]
# 70 registros: 60 de 2024 (páginas de 20: 20, 20, 20 e a vazia) e 10 de 2023
PRODUCAO = [
    {
        "ano": 2024 if i < 60 else 2023,
        "uf": ["SP", "RJ", "MG", "SP", "BA"][i % 5],
        "sexo": [None, "F", "M"][i % 3],
        "procedimento": f"Procedimento {'ABCD'[i % 4]}",
        "quantidade": str(i % 7 + 1),
        "valor": VALORES[i % len(VALORES)],
        "municipio": {"codigo": f"35{i % 3:04d}"},
    }
    for i in range(70)
]


class ApiProducao:
    """Imita a API oficial: offset = número da página. Com ``lenta``, trava a partir da 2ª página."""

    def __init__(self, lenta: bool = False):
        self.lenta = lenta

    async def __call__(self, request: httpx.Request) -> httpx.Response:
        if request.url.path == "/spec.json":
            return httpx.Response(200, json=SPEC)
        if request.url.path != "/sia/producao":
            return httpx.Response(404, json={"erro": "não encontrado"})
        q = request.url.params
        limit, offset = int(q.get("limit", 20)), int(q.get("offset", 0))
        if self.lenta and offset > 0:
            await asyncio.sleep(60)
        dados = [r for r in PRODUCAO if str(r["ano"]) == q.get("ano")]
        return httpx.Response(200, json={"producao": dados[offset * limit:(offset + 1) * limit]})


def _num(texto):
    """Conversão independente da usada pelo app: "1.234,56", "10,5", "100.25", "7" ou None."""
    if texto is None:
        return 0.0
    if re.fullmatch(r"\d{1,3}(\.\d{3})+(,\d+)?", texto):
        return float(texto.replace(".", "").replace(",", "."))
    return float(texto.replace(",", "."))


@pytest.fixture
def pasta(tmp_path):
    return tmp_path / "dados"


@pytest.fixture
def abrir(config, pasta):
    config.pasta_dados = pasta

    def _abrir(api=None) -> TestClient:
        return TestClient(criar_app(config, transport=httpx.MockTransport(api or ApiProducao())))

    return _abrir


def esperar(c: TestClient, download_id: str, condicao=lambda d: d["status"] in FINAIS, timeout: float = 15.0) -> dict:
    fim = time.monotonic() + timeout
    while True:
        d = c.get(f"/api/downloads/{download_id}").json()
        if condicao(d):
            return d
        if time.monotonic() > fim:
            raise AssertionError(f"tempo esgotado esperando o download: {d}")
        time.sleep(0.02)


def test_criar_app_registra_as_rotas_do_app_local(abrir):
    with abrir() as c:
        caminhos = set(c.get("/openapi.json").json()["paths"])
        assert {"/api/downloads", "/api/downloads/{download_id}", "/api/local/bases"} <= caminhos
        assert "/api/local/bases/{nome}/categorias/{coluna}" in caminhos
        assert c.get("/api/downloads").json() == {"downloads": []}
        assert c.get("/health").json()["app_local"] is True


def test_fluxo_completo_baixar_filtrar_agregar_exportar(abrir, pasta):
    base_2024 = [r for r in PRODUCAO if r["ano"] == 2024]
    with abrir() as c:
        # 1. baixar
        r = c.post("/api/downloads", json={"dataset": "sia-producao", "filtros": {"ano": 2024}})
        assert r.status_code == 201, r.text
        d = esperar(c, r.json()["id"])
        assert d["status"] == "concluido", d
        assert d["nome"] == "sia-producao__ano-2024"
        assert d["registros"] == 60 and d["paginas"] == 4
        nome = d["nome"]
        assert c.get("/api/downloads").json()["downloads"][0]["id"] == d["id"]

        # 2. listar
        lista = c.get("/api/local/bases").json()
        assert lista["pasta"] == str(pasta)
        [b] = lista["bases"]
        assert b["nome"] == nome and b["dataset"] == "sia-producao" and b["registros"] == 60
        assert b["filtros"] == {"ano": 2024} and b["titulo"] == "Produção ambulatorial (sintética)"
        assert b["tamanho_bytes"] > 0 and b["paginas"] == 4
        assert set(b["colunas"]) == {"ano", "uf", "sexo", "procedimento", "quantidade", "valor", "municipio.codigo"}

        # 3. colunas
        colunas = c.get(f"/api/local/bases/{nome}/colunas").json()
        assert colunas["registros"] == 60
        info = {col["nome"]: col for col in colunas["colunas"]}
        assert info["quantidade"]["numerica"] and info["valor"]["numerica"]
        assert not info["uf"]["numerica"] and not info["procedimento"]["numerica"]
        assert info["uf"]["distintos"] == 4 and info["sexo"]["nulos"] == 20
        assert info["valor"]["nulos"] == 12  # 1 em cada 5

        # 4. categorias (com filtros em cascata)
        cat = c.get(f"/api/local/bases/{nome}/categorias/uf").json()
        por_uf = Counter(r["uf"] for r in base_2024)
        assert {v["valor"]: v["registros"] for v in cat["valores"]} == dict(por_uf)
        assert cat["valores"][0] == {"valor": "SP", "registros": por_uf["SP"]}
        assert cat["total_distintos"] == 4 and cat["truncado"] is False

        filtros_sexo = json.dumps({"sexo": ["F"], "uf": ["MG"]})  # o filtro da própria coluna é ignorado
        cat = c.get(f"/api/local/bases/{nome}/categorias/uf", params={"filtros": filtros_sexo}).json()
        esperado = Counter(r["uf"] for r in base_2024 if r["sexo"] == "F")
        assert {v["valor"]: v["registros"] for v in cat["valores"]} == dict(esperado)

        cat = c.get(f"/api/local/bases/{nome}/categorias/sexo", params={"filtros": json.dumps({"uf": ["SP", "RJ"]})}).json()
        esperado = Counter(r["sexo"] for r in base_2024 if r["uf"] in ("SP", "RJ"))
        assert {v["valor"]: v["registros"] for v in cat["valores"]} == dict(esperado)
        assert cat["valores"][-1]["valor"] is None  # o nulo vem por último

        cat = c.get(f"/api/local/bases/{nome}/categorias/procedimento", params={"busca": "procedimento b", "limite": 1}).json()
        assert [v["valor"] for v in cat["valores"]] == ["Procedimento B"] and cat["truncado"] is False

        # 5. consulta com filtros e agregação
        filtros = {"uf": ["SP", "RJ"], "sexo": ["F", None]}
        corpo = {"filtros": filtros, "colunas": None, "agrupar_por": ["procedimento"], "somar": ["valor", "quantidade"], "limite": 5}
        res = c.post(f"/api/local/bases/{nome}/consulta", json=corpo).json()
        filtrados = [r for r in base_2024 if r["uf"] in ("SP", "RJ") and r["sexo"] in ("F", None)]
        assert res["total"] == len(filtrados) and len(res["linhas"]) == 5
        assert all(linha["uf"] in ("SP", "RJ") and linha["sexo"] in ("F", None) for linha in res["linhas"])

        grupos = defaultdict(lambda: [0, 0.0, 0.0])
        for r in filtrados:
            g = grupos[r["procedimento"]]
            g[0] += 1
            g[1] += _num(r["valor"])
            g[2] += _num(r["quantidade"])
        ag = res["agregado"]
        assert ag["colunas"] == ["procedimento", "registros", "soma_valor", "soma_quantidade"]
        assert ag["grupos"] == len(grupos) == len(ag["linhas"])
        somas = [linha["soma_valor"] for linha in ag["linhas"]]
        assert somas == sorted(somas, reverse=True)  # ordenado pela 1ª soma
        for linha in ag["linhas"]:
            n, valor, qtd = grupos[linha["procedimento"]]
            assert linha["registros"] == n
            assert linha["soma_valor"] == pytest.approx(valor)
            assert linha["soma_quantidade"] == qtd and isinstance(linha["soma_quantidade"], int)  # sem ".0"
        assert ag["totais"]["registros"] == len(filtrados)
        assert ag["totais"]["soma_valor"] == pytest.approx(sum(g[1] for g in grupos.values()))
        assert ag["totais"]["soma_quantidade"] == sum(g[2] for g in grupos.values())

        # 6. exportar para Excel
        corpo_exp = {k: v for k, v in corpo.items() if k != "limite"}
        r = c.post(f"/api/local/bases/{nome}/exportar", json={**corpo_exp, "formato": "xlsx"})
        assert r.status_code == 200
        assert r.headers["content-disposition"].startswith("attachment;")
        assert f'filename="{nome}__' in r.headers["content-disposition"]
        salvo = pasta / "exportacoes" / unquote(r.headers["x-arquivo-salvo"]).rsplit("/", 1)[-1]
        assert salvo.is_file() and salvo.read_bytes() == r.content
        wb = load_workbook(io.BytesIO(r.content), read_only=True)
        assert wb.sheetnames == ["resumo", "dados", "consulta"]
        resumo = [list(linha) for linha in wb["resumo"].iter_rows(values_only=True)]
        assert resumo[0] == ag["colunas"]
        assert [linha[0] for linha in resumo[1:-1]] == [linha["procedimento"] for linha in ag["linhas"]]
        assert resumo[-1][:2] == ["TOTAL", len(filtrados)]
        assert resumo[-1][2] == pytest.approx(ag["totais"]["soma_valor"])
        dados = list(wb["dados"].iter_rows(values_only=True))
        assert len(dados) == len(filtrados) + 1 and "municipio.codigo" in dados[0]
        consulta = {linha[0]: linha[1] for linha in wb["consulta"].iter_rows(values_only=True)}
        assert json.loads(consulta["Filtros"]) == filtros

        # 7. exportar para CSV (o resumo, com ";" e BOM)
        r = c.post(f"/api/local/bases/{nome}/exportar", json={**corpo_exp, "formato": "csv"})
        assert r.status_code == 200 and r.content.startswith(b"\xef\xbb\xbf")
        linhas_csv = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig")), delimiter=";"))
        assert linhas_csv[0] == ag["colunas"] and linhas_csv[-1][:2] == ["TOTAL", str(len(filtrados))]
        assert len(list((pasta / "exportacoes").iterdir())) == 2

        # 8. excluir
        assert c.delete(f"/api/local/bases/{nome}").json() == {"removida": nome}
        assert c.get("/api/local/bases").json()["bases"] == []
        assert c.get(f"/api/local/bases/{nome}/colunas").status_code == 404


def test_encerrar_o_app_cancela_downloads_e_apaga_parciais(abrir, pasta):
    with abrir(ApiProducao(lenta=True)) as c:
        r = c.post("/api/downloads", json={"dataset": "sia-producao", "filtros": {"ano": 2024}})
        assert r.status_code == 201
        d = esperar(c, r.json()["id"], lambda d: d["registros"] > 0)
        assert d["status"] == "baixando"
        assert any(p.name.endswith(".parcial") for p in (pasta / "bases").iterdir())
        gerenciador, cliente = c.app.state.downloads, c.app.state.cliente
        fechar_original, ao_fechar = cliente.fechar, {}

        async def fechar():  # registra o estado no momento em que o cliente HTTP é fechado
            ao_fechar["status"] = gerenciador.obter(d["id"])["status"]
            ao_fechar["arquivos"] = sorted(p.name for p in (pasta / "bases").iterdir())
            await fechar_original()

        cliente.fechar = fechar
    # ao sair do "with", o ciclo de vida do app cancela o download ANTES de fechar o cliente HTTP
    assert ao_fechar == {"status": "cancelado", "arquivos": []}
    assert gerenciador.obter(d["id"])["status"] == "cancelado"
    assert list((pasta / "bases").iterdir()) == []


def test_download_invalido_nao_cria_nada(abrir, pasta):
    with abrir() as c:
        assert c.post("/api/downloads", json={"dataset": "nao-existe"}).status_code == 404
        r = c.post("/api/downloads", json={"dataset": "sia-producao", "filtros": {}})
        assert r.status_code == 422 and "ano" in r.json()["erro"]
        assert c.get("/api/local/bases").json()["bases"] == []
