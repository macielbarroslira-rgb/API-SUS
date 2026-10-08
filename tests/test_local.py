"""Consultas sobre as bases baixadas para o computador (app/local.py e app/rotas_local.py)."""

import csv
import io
import json
import math
import re

import duckdb
import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import local
from app.agregacao import _limpo, numero
from app.main import criar_app
from app.rotas_local import router
from tests.conftest import ApiFalsa

COLUNAS = ["uf", "municipio", "procedimento", "qtd", "valor", "endereco.cep"]
PRODUCAO = [
    dict(zip(COLUNAS, linha))
    for linha in [
        ("SP", "São Paulo", "Consulta", "10", "1.234,56", "01000"),
        ("SP", "Campinas", "Consulta", "5", "10,5", "13000"),
        ("SP", "São Paulo", "Exame", "3", "100.25", None),
        ("RJ", "Rio de Janeiro", "Consulta", "7", "2", "20000"),
        ("RJ", "Niterói", "Exame", "", "abc", "24000"),
        (None, "Desconhecido", "Exame", "1", None, None),
        ("MG", "Olhos D'Água \"x\"", "=1+1", "2", "1.000", "39000"),
    ]
]


def criar_base(pasta, nome, registros, meta=None):
    """Grava bases/<nome>.parquet (todas as colunas VARCHAR) e bases/<nome>.json, como o app faz."""
    pb = pasta / "bases"
    pb.mkdir(parents=True, exist_ok=True)
    colunas = list(dict.fromkeys(k for r in registros for k in r))
    con = duckdb.connect(":memory:")
    try:
        con.execute("CREATE TABLE t (" + ", ".join(f"{local.citar(c)} VARCHAR" for c in colunas) + ")")
        con.executemany(
            "INSERT INTO t VALUES (" + ", ".join("?" for _ in colunas) + ")",
            [[r.get(c) for c in colunas] for r in registros],
        )
        con.table("t").write_parquet(str(pb / f"{nome}.parquet"))
    finally:
        con.close()
    metadados = {
        "nome": nome,
        "dataset": "sia-producao",
        "titulo": "Produção ambulatorial",
        "filtros": {"nu_ano": 2024},
        "registros": len(registros),
        "colunas": colunas,
        "baixado_em": "2026-10-01T12:00:00+00:00",
        "url_origem": "https://api.teste/sia/producao",
        "avisos": [],
        "paginas": 1,
        **(meta or {}),
    }
    (pb / f"{nome}.json").write_text(json.dumps(metadados, ensure_ascii=False), encoding="utf-8")
    return pb / f"{nome}.parquet"


@pytest.fixture
def pasta(tmp_path):
    return tmp_path / "dados"


@pytest.fixture
def cli(config, pasta):
    """App de teste com o router das bases locais (e a base "producao" já baixada)."""
    criar_base(pasta, "producao", PRODUCAO)
    config.pasta_dados = pasta
    app = criar_app(config, transport=httpx.MockTransport(ApiFalsa()))
    if not any(getattr(r, "path", None) == "/api/local/bases" for r in app.routes):
        app.include_router(router)
    with TestClient(app) as c:
        yield c


def consulta(cli, nome="producao", **corpo):
    return cli.post(f"/api/local/bases/{nome}/consulta", json=corpo)


# --------------------------------------------------------------------------- #
# Listar e excluir
# --------------------------------------------------------------------------- #


def test_lista_bases_mais_recentes_primeiro(cli, pasta):
    criar_base(pasta, "nova", PRODUCAO[:2], {"baixado_em": "2026-10-05T08:00:00+00:00"})
    criar_base(pasta, "velha", PRODUCAO[:1], {"baixado_em": "2025-01-01T00:00:00+00:00"})
    pb = pasta / "bases"
    (pb / "parcial.json.parcial").write_text("{}")  # download em andamento
    (pb / "Invalida.json").write_text("{}")  # nome fora do padrão
    (pb / "sem-dados.json").write_text("{}")  # metadados sem o parquet

    r = cli.get("/api/local/bases").json()
    assert r["pasta"] == str(pasta)
    assert [b["nome"] for b in r["bases"]] == ["nova", "producao", "velha"]
    nova = r["bases"][0]
    assert nova["registros"] == 2 and nova["dataset"] == "sia-producao" and nova["filtros"] == {"nu_ano": 2024}
    tamanho = (pb / "nova.parquet").stat().st_size + (pb / "nova.json").stat().st_size
    assert nova["tamanho_bytes"] == tamanho
    assert (pasta / "exportacoes").is_dir()


def test_lista_pasta_vazia_cria_subpastas(tmp_path):
    pasta = tmp_path / "nao-existe"
    assert local.listar_bases(pasta) == []
    assert (pasta / "bases").is_dir() and (pasta / "exportacoes").is_dir()


def test_lista_tolera_metadados_ilegiveis(pasta):
    criar_base(pasta, "quebrada", PRODUCAO)
    (pasta / "bases" / "quebrada.json").write_text("{nao é json", encoding="utf-8")
    [base] = local.listar_bases(pasta)
    assert base["nome"] == "quebrada" and base["registros"] == len(PRODUCAO)
    assert any("ilegíveis" in a for a in base["avisos"])


def test_exclui_base(cli, pasta):
    assert cli.delete("/api/local/bases/producao").json() == {"removida": "producao"}
    assert not (pasta / "bases" / "producao.parquet").exists()
    assert not (pasta / "bases" / "producao.json").exists()
    assert cli.get("/api/local/bases").json()["bases"] == []
    assert cli.delete("/api/local/bases/producao").status_code == 404


# --------------------------------------------------------------------------- #
# Colunas e categorias
# --------------------------------------------------------------------------- #


def test_colunas(cli):
    r = cli.get("/api/local/bases/producao/colunas").json()
    assert r["nome"] == "producao" and r["registros"] == 7
    cols = {c["nome"]: c for c in r["colunas"]}
    assert list(cols) == ["uf", "municipio", "procedimento", "qtd", "valor", "endereco.cep"]
    assert cols["uf"] == {"nome": "uf", "distintos": 3, "nulos": 1, "numerica": False, "exemplo": "SP"}
    assert cols["qtd"]["numerica"] is True  # o vazio não conta contra
    assert cols["qtd"]["distintos"] == 7 and cols["qtd"]["nulos"] == 0
    assert cols["valor"]["numerica"] is False  # "abc" em 6 preenchidos (83%)
    assert cols["endereco.cep"]["nulos"] == 2


@pytest.mark.parametrize("ruins, numerica", [(0, True), (1, True), (2, False)])
def test_coluna_numerica_com_95_por_cento(pasta, ruins, numerica):
    valores = ["texto"] * ruins + [str(i) for i in range(20 - ruins)] + [None, "  "]
    criar_base(pasta, "limiar", [{"v": v} for v in valores])
    [col] = local.colunas(pasta, "limiar")["colunas"]
    assert col["numerica"] is numerica


def test_coluna_so_vazia_nao_e_numerica(pasta):
    criar_base(pasta, "vazia", [{"v": None}, {"v": ""}])
    [col] = local.colunas(pasta, "vazia")["colunas"]
    assert col["numerica"] is False and col["nulos"] == 1 and col["exemplo"] == ""


def test_distintos_aproximados_em_base_grande(pasta, monkeypatch):
    monkeypatch.setattr(local, "DISTINTOS_EXATOS_ATE", 3)
    criar_base(pasta, "grande", PRODUCAO)
    r = local.colunas(pasta, "grande")
    assert r["distintos_aproximados"] is True
    assert {c["nome"]: c["distintos"] for c in r["colunas"]}["uf"] == 3


def test_categorias_contagens_e_nulo(cli):
    r = cli.get("/api/local/bases/producao/categorias/uf").json()
    assert r["coluna"] == "uf" and r["total_distintos"] == 4 and r["truncado"] is False
    assert r["valores"] == [
        {"valor": "SP", "registros": 3},
        {"valor": "RJ", "registros": 2},
        {"valor": "MG", "registros": 1},
        {"valor": None, "registros": 1},
    ]


def test_categorias_busca_sem_acento_e_sem_curingas(cli):
    url = "/api/local/bases/producao/categorias/municipio"
    r = cli.get(url, params={"busca": "SAO"}).json()
    assert r["valores"] == [{"valor": "São Paulo", "registros": 2}]
    assert cli.get(url, params={"busca": "niteroi"}).json()["total_distintos"] == 1
    # % e _ são texto comum (não curingas do LIKE)
    assert cli.get(url, params={"busca": "%"}).json()["valores"] == []
    assert cli.get(url, params={"busca": "_"}).json()["valores"] == []


def test_categorias_em_cascata_ignoram_a_propria_coluna(cli):
    filtros = json.dumps({"uf": ["SP"], "municipio": ["Campinas"]})
    url = "/api/local/bases/producao/categorias/"
    # municípios: o filtro de município é ignorado, o de UF vale
    r = cli.get(url + "municipio", params={"filtros": filtros}).json()
    assert [v["valor"] for v in r["valores"]] == ["São Paulo", "Campinas"]
    # procedimentos: valem os dois filtros
    r = cli.get(url + "procedimento", params={"filtros": filtros}).json()
    assert r["valores"] == [{"valor": "Consulta", "registros": 1}]
    # null na lista = valor vazio
    r = cli.get(url + "municipio", params={"filtros": json.dumps({"uf": [None]})}).json()
    assert r["valores"] == [{"valor": "Desconhecido", "registros": 1}]


def test_categorias_limite_e_truncado(cli):
    url = "/api/local/bases/producao/categorias/municipio"
    r = cli.get(url, params={"limite": 2}).json()
    assert len(r["valores"]) == 2 and r["total_distintos"] == 6 and r["truncado"] is True
    assert r["valores"][0] == {"valor": "São Paulo", "registros": 2}
    assert cli.get(url, params={"limite": 0}).status_code == 422
    assert cli.get(url, params={"limite": local.LIMITE_CATEGORIAS_MAX + 1}).status_code == 422


def test_categorias_de_coluna_com_ponto(cli):
    r = cli.get("/api/local/bases/producao/categorias/endereco.cep", params={"limite": 1}).json()
    assert r["valores"] == [{"valor": None, "registros": 2}] and r["total_distintos"] == 6


def test_404_para_base_e_coluna_inexistentes(cli):
    assert cli.get("/api/local/bases/nao-existe/colunas").status_code == 404
    assert cli.get("/api/local/bases/nao-existe/categorias/uf").status_code == 404
    assert cli.get("/api/local/bases/producao/categorias/nao_existe").status_code == 404
    assert consulta(cli, "nao-existe").status_code == 404
    assert consulta(cli, colunas=["nao_existe"]).status_code == 404
    assert consulta(cli, filtros={"nao_existe": ["x"]}).status_code == 404
    assert consulta(cli, agrupar_por=["nao_existe"]).status_code == 404
    assert consulta(cli, somar=["nao_existe"]).status_code == 404
    r = cli.post("/api/local/bases/nao-existe/exportar", json={"formato": "csv"})
    assert r.status_code == 404
    assert "nao-existe" in r.json()["detail"]


# --------------------------------------------------------------------------- #
# Consulta
# --------------------------------------------------------------------------- #


def test_consulta_sem_filtros(cli):
    r = consulta(cli).json()
    assert r["total"] == 7 and len(r["linhas"]) == 7 and r["agregado"] is None
    assert r["colunas"] == ["uf", "municipio", "procedimento", "qtd", "valor", "endereco.cep"]
    assert r["linhas"][0]["endereco.cep"] == "01000"


def test_consulta_filtros_in_com_and_e_nulo(cli):
    r = consulta(cli, filtros={"uf": ["SP", "RJ"], "procedimento": ["Exame"]}).json()
    assert r["total"] == 2
    assert sorted(l["municipio"] for l in r["linhas"]) == ["Niterói", "São Paulo"]
    r = consulta(cli, filtros={"uf": ["MG", None]}).json()
    assert sorted(l["municipio"] for l in r["linhas"]) == ["Desconhecido", "Olhos D'Água \"x\""]
    # lista vazia = sem filtro naquela coluna; valor solto vale como lista de um
    assert consulta(cli, filtros={"uf": []}).json()["total"] == 7
    assert consulta(cli, filtros={"uf": "RJ"}).json()["total"] == 2


def test_consulta_colunas_e_limite(cli):
    r = consulta(cli, colunas=["municipio", "uf"], limite=2).json()
    assert r["total"] == 7 and r["colunas"] == ["municipio", "uf"]
    assert r["linhas"] == [{"municipio": "São Paulo", "uf": "SP"}, {"municipio": "Campinas", "uf": "SP"}]
    assert consulta(cli, limite=0).json()["linhas"] == []
    assert consulta(cli, limite=-1).status_code == 422
    assert consulta(cli, limite=local.LIMITE_CONSULTA_MAX + 1).status_code == 422


def test_consulta_agrupa_e_soma(cli):
    r = consulta(cli, agrupar_por=["uf"], somar=["valor", "qtd"], limite=0).json()
    ag = r["agregado"]
    assert ag["agrupar_por"] == ["uf"] and ag["somar"] == ["valor", "qtd"]
    assert ag["colunas"] == ["uf", "registros", "soma_valor", "soma_qtd"]
    assert ag["grupos"] == 4
    # ordenado pela 1ª soma (decrescente); "1.234,56" = 1234.56, "10,5" = 10.5, "1.000" = 1000
    assert ag["linhas"] == [
        {"uf": "SP", "registros": 3, "soma_valor": 1345.31, "soma_qtd": 18},
        {"uf": "MG", "registros": 1, "soma_valor": 1000, "soma_qtd": 2},
        {"uf": "RJ", "registros": 2, "soma_valor": 2, "soma_qtd": 7},
        {"uf": None, "registros": 1, "soma_valor": 0, "soma_qtd": 1},
    ]
    assert ag["totais"] == {"registros": 7, "soma_valor": 2347.31, "soma_qtd": 28}
    assert isinstance(ag["totais"]["soma_qtd"], int)  # inteiro sem ".0"


def test_consulta_agrupa_sem_soma_ordena_por_registros(cli):
    ag = consulta(cli, agrupar_por=["procedimento", "uf"], filtros={"uf": ["SP", "RJ"]}).json()["agregado"]
    assert ag["colunas"] == ["procedimento", "uf", "registros"]
    assert [l["registros"] for l in ag["linhas"]] == [2, 1, 1, 1]
    assert ag["linhas"][0] == {"procedimento": "Consulta", "uf": "SP", "registros": 2}
    assert ag["totais"] == {"registros": 5}


def test_consulta_so_soma_e_sem_registros(cli):
    ag = consulta(cli, somar=["qtd"]).json()["agregado"]
    assert ag["colunas"] == ["registros", "soma_qtd"] and ag["grupos"] == 1
    assert ag["linhas"] == [{"registros": 7, "soma_qtd": 28}] and ag["totais"] == {"registros": 7, "soma_qtd": 28}
    vazio = consulta(cli, somar=["qtd"], agrupar_por=["uf"], filtros={"uf": ["XX"]}).json()
    assert vazio["total"] == 0
    assert vazio["agregado"]["linhas"] == [] and vazio["agregado"]["grupos"] == 0
    assert vazio["agregado"]["totais"] == {"registros": 0, "soma_qtd": 0}
    assert consulta(cli, somar=["qtd"], filtros={"uf": ["XX"]}).json()["agregado"]["linhas"] == []


def test_consulta_limita_grupos_na_resposta(cli, monkeypatch):
    monkeypatch.setattr(local, "MAX_GRUPOS_RESPOSTA", 2)
    ag = consulta(cli, agrupar_por=["municipio"]).json()["agregado"]
    assert ag["grupos"] == 6 and len(ag["linhas"]) == 2 and ag["totais"]["registros"] == 7


def test_consulta_conflito_de_nomes_no_resumo(pasta):
    criar_base(pasta, "conflito", [{"registros": "1", "x": "2"}])
    with pytest.raises(local.ErroConsulta):
        local.consultar(pasta, "conflito", agrupar_por=["registros"])


SOMAS_DE_BORDA = [
    "10", "-10", "007", "10.5", "10,5", "-0,25", "1.234,56", "-1.234,5", "1.234", "1.234.567", "1.234.567,89",
    "12.34", "1,234.56", "1.23,4", "12,", ",5", "", "   ", " 12 ", "\t3\n", "\xa04　", "abc", "R$ 10",
    "1e3", "-2.5E-2", "+5", ".5", "5.", "1_000", "1__0", "_1", "0x10", "nan", "inf", "-Infinity", "1e400",
    "1 000", "10%",
]


def test_soma_sql_igual_a_numero(pasta):
    """A soma feita no DuckDB segue a mesma regra de app.agregacao.numero (em casos de borda)."""
    criar_base(pasta, "bordas", [{"v": v} for v in SOMAS_DE_BORDA])
    ag = local.consultar(pasta, "bordas", agrupar_por=["v"], somar=["v"], limite=0)["agregado"]
    somas = {l["v"]: l["soma_v"] for l in ag["linhas"]}
    esperado_total = 0.0
    for v in SOMAS_DE_BORDA:
        n = numero(v)
        n = n if n is not None and math.isfinite(n) else None  # nan/inf não entram na soma
        esperado_total += n or 0.0
        assert somas[v] == (_limpo(n) if n is not None else 0), v
    assert ag["totais"]["soma_v"] == _limpo(esperado_total)
    # a coluna é numérica só se >= 95% dos preenchidos forem números
    [col] = local.colunas(pasta, "bordas")["colunas"]
    assert col["numerica"] is False


def test_soma_sql_espacos_iguais_ao_strip_do_python():
    espacos = "".join(chr(i) for i in range(0x110000) if chr(i).isspace())
    assert local._ESPACOS == espacos


# --------------------------------------------------------------------------- #
# Exportação
# --------------------------------------------------------------------------- #


def _abrir_xlsx(conteudo):
    return load_workbook(io.BytesIO(conteudo))


def test_exporta_xlsx_com_resumo(cli, pasta):
    corpo = {"filtros": {"uf": ["SP", "RJ"]}, "agrupar_por": ["uf"], "somar": ["qtd"], "formato": "xlsx"}
    r = cli.post("/api/local/bases/producao/exportar", json=corpo)
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    disposicao = r.headers["content-disposition"]
    assert disposicao.startswith("attachment") and re.search(r'filename="producao__\d{8}-\d{6}\.xlsx"', disposicao)

    salvos = list((pasta / "exportacoes").iterdir())
    assert len(salvos) == 1 and salvos[0].read_bytes() == r.content
    assert re.fullmatch(r"producao__\d{8}-\d{6}\.xlsx", salvos[0].name)

    wb = _abrir_xlsx(r.content)
    assert wb.sheetnames == ["resumo", "dados", "consulta"]
    resumo = [list(l) for l in wb["resumo"].iter_rows(values_only=True)]
    assert resumo == [["uf", "registros", "soma_qtd"], ["SP", 3, 18], ["RJ", 2, 7], ["TOTAL", 5, 25]]
    ultima = list(wb["resumo"].iter_rows())[-1]
    assert all(c.font.bold for c in ultima)
    dados = list(wb["dados"].iter_rows(values_only=True))
    assert dados[0] == ("uf", "municipio", "procedimento", "qtd", "valor", "endereco.cep") and len(dados) == 6
    info = {l[0]: l[1] for l in wb["consulta"].iter_rows(values_only=True)}
    assert info["Base"] == "producao" and info["Base de origem"] == "sia-producao"
    assert json.loads(info["Filtros"]) == {"uf": ["SP", "RJ"]}
    assert info["Registros na base"] == 7 and info["Registros filtrados"] == 5 and info["Grupos"] == 2
    assert "Aviso" not in info


def test_exporta_xlsx_sem_resumo_e_texto_nao_vira_formula(cli):
    corpo = {"filtros": {"uf": ["MG"]}, "colunas": ["municipio", "procedimento"]}
    r = cli.post("/api/local/bases/producao/exportar", json=corpo)
    wb = _abrir_xlsx(r.content)
    assert wb.sheetnames == ["dados", "consulta"]
    _, linha = list(wb["dados"].iter_rows())
    assert linha[0].value == "Olhos D'Água \"x\""
    assert linha[1].value == "=1+1" and linha[1].data_type == "s"  # texto, não fórmula


def test_exporta_xlsx_corta_no_limite_do_excel(pasta, monkeypatch):
    monkeypatch.setattr(local, "LINHAS_MAX_XLSX", 3)
    criar_base(pasta, "corte", PRODUCAO)
    caminho, sugerido = local.exportar(pasta, "corte", "xlsx")
    assert caminho.parent == pasta / "exportacoes" and caminho.name == sugerido
    wb = load_workbook(caminho)
    assert wb["dados"].max_row == 4  # cabeçalho + 3
    avisos = [l[1] for l in wb["consulta"].iter_rows(values_only=True) if l[0] == "Aviso"]
    assert len(avisos) == 1 and "7 registros" in avisos[0] and "CSV" in avisos[0]


def test_exporta_csv_resumo_com_total(cli):
    corpo = {"agrupar_por": ["uf"], "somar": ["valor"], "formato": "csv"}
    r = cli.post("/api/local/bases/producao/exportar", json=corpo)
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
    assert r.content.startswith(b"\xef\xbb\xbf")
    linhas = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig")), delimiter=";"))
    assert linhas[0] == ["uf", "registros", "soma_valor"]
    assert linhas[1] == ["SP", "3", "1345.31"]
    assert linhas[4] == ["", "1", "0"]  # grupo dos nulos
    assert linhas[-1] == ["TOTAL", "7", "2347.31"]


def test_exporta_csv_dados(cli, pasta):
    corpo = {"filtros": {"uf": ["MG"]}, "formato": "csv"}
    r = cli.post("/api/local/bases/producao/exportar", json=corpo)
    texto = r.content.decode("utf-8-sig")
    linhas = list(csv.reader(io.StringIO(texto), delimiter=";"))
    assert linhas == [
        ["uf", "municipio", "procedimento", "qtd", "valor", "endereco.cep"],
        ["MG", "Olhos D'Água \"x\"", "=1+1", "2", "1.000", "39000"],
    ]
    assert ";" in texto.splitlines()[0]
    assert re.search(r'filename="producao__\d{8}-\d{6}\.csv"', r.headers["content-disposition"])
    assert not list((pasta / "exportacoes").glob("*.parcial"))


def test_duas_exportacoes_no_mesmo_segundo_nao_se_sobrescrevem(pasta):
    criar_base(pasta, "dupla", PRODUCAO)
    a, _ = local.exportar(pasta, "dupla", "csv")
    b, _ = local.exportar(pasta, "dupla", "csv")
    assert a != b and a.exists() and b.exists()


def test_exporta_formato_invalido(cli, pasta):
    assert cli.post("/api/local/bases/producao/exportar", json={"formato": "pdf"}).status_code == 422
    with pytest.raises(local.ErroConsulta):
        local.exportar(pasta, "producao", "pdf")


# --------------------------------------------------------------------------- #
# Tentativas de injeção e entradas inválidas
# --------------------------------------------------------------------------- #

COLUNAS_MALICIOSAS = [
    'uf"; DROP TABLE base; --',
    "uf' OR '1'='1",
    "uf) OR (1=1",
    '"uf"',
    "UF",  # o DuckDB não diferencia maiúsculas, mas a coluna precisa existir exatamente
    "*",
    "",
]


@pytest.mark.parametrize("coluna", COLUNAS_MALICIOSAS)
def test_colunas_maliciosas_sao_404(cli, coluna):
    assert consulta(cli, filtros={coluna: ["SP"]}).status_code == 404
    assert consulta(cli, colunas=[coluna]).status_code == 404
    assert consulta(cli, agrupar_por=[coluna]).status_code == 404
    assert consulta(cli, somar=[coluna]).status_code == 404
    r = cli.post("/api/local/bases/producao/exportar", json={"colunas": [coluna], "formato": "csv"})
    assert r.status_code == 404
    if coluna:
        assert cli.get("/api/local/bases/producao/categorias/" + coluna.replace("%", "%25")).status_code == 404
    # nada foi apagado
    assert consulta(cli).json()["total"] == 7


@pytest.mark.parametrize(
    "valor", ["SP' OR '1'='1", "'; DROP TABLE base; --", "SP\" OR \"1\"=\"1", "%", "_", "SP' --"]
)
def test_valores_maliciosos_sao_texto_literal(cli, valor):
    r = consulta(cli, filtros={"uf": [valor]}).json()
    assert r["total"] == 0 and r["linhas"] == []
    r = cli.get(
        "/api/local/bases/producao/categorias/municipio", params={"filtros": json.dumps({"uf": [valor]})}
    ).json()
    assert r["valores"] == []
    assert consulta(cli).json()["total"] == 7


def test_valor_com_aspas_casa_literalmente(cli):
    r = consulta(cli, filtros={"municipio": ["Olhos D'Água \"x\""]}).json()
    assert r["total"] == 1 and r["linhas"][0]["uf"] == "MG"
    r = cli.get("/api/local/bases/producao/categorias/municipio", params={"busca": "D'Água \"x"}).json()
    assert r["valores"] == [{"valor": "Olhos D'Água \"x\"", "registros": 1}]


def test_coluna_com_nome_estranho_funciona_citada(pasta):
    estranha = 'nome "estranho"; DROP TABLE base; --'
    criar_base(pasta, "estranha", [{estranha: "a", "n": "1"}, {estranha: "a", "n": "2"}, {estranha: "b", "n": "x"}])
    r = local.consultar(pasta, "estranha", filtros={estranha: ["a"]}, agrupar_por=[estranha], somar=["n"])
    assert r["total"] == 2 and r["linhas"][0][estranha] == "a"
    assert r["agregado"]["linhas"] == [{estranha: "a", "registros": 2, "soma_n": 3}]
    cat = local.categorias(pasta, "estranha", estranha)
    assert cat["valores"] == [{"valor": "a", "registros": 2}, {"valor": "b", "registros": 1}]
    cols = local.colunas(pasta, "estranha")["colunas"]
    assert [c["nome"] for c in cols] == [estranha, "n"]


@pytest.mark.parametrize("nome", ["../x", "../../etc/passwd", "..", "/tmp/x", "Producao", "a.b", "a b", "", "x" * 81])
def test_nomes_de_base_invalidos(pasta, nome):
    criar_base(pasta, "producao", PRODUCAO)
    for funcao in (local.colunas, local.consultar, local.excluir_base):
        with pytest.raises(local.BaseNaoEncontrada):
            funcao(pasta, nome)
    with pytest.raises(local.BaseNaoEncontrada):
        local.categorias(pasta, nome, "uf")
    with pytest.raises(local.BaseNaoEncontrada):
        local.exportar(pasta, nome, "csv")


def test_nome_de_base_fora_da_pasta_nao_e_lido(cli, pasta, tmp_path):
    # uma base válida fora de bases/ não pode ser alcançada por caminhos relativos
    criar_base(tmp_path, "fora", PRODUCAO)
    for url in (
        "/api/local/bases/%2E%2E/colunas",
        "/api/local/bases/..%2Ffora/colunas",
        "/api/local/bases/..%2F..%2Fbases%2Ffora/colunas",
    ):
        assert cli.get(url).status_code == 404, url
    assert cli.delete("/api/local/bases/%2E%2E").status_code in (404, 405)
    assert cli.delete("/api/local/bases/..%2Fbases%2Fproducao").status_code in (404, 405)
    assert (pasta / "bases" / "producao.parquet").exists()


@pytest.mark.parametrize(
    "filtros",
    ["{nao json", "[1, 2]", '"texto"', '{"uf": [{"a": 1}]}', '{"uf": [[1]]}'],
)
def test_filtros_invalidos_na_query_sao_422(cli, filtros):
    r = cli.get("/api/local/bases/producao/categorias/municipio", params={"filtros": filtros})
    assert r.status_code == 422 and "erro" in r.json()


def test_corpo_invalido_e_422(cli):
    assert consulta(cli, filtros={"uf": [{"x": 1}]}).status_code == 422
    assert consulta(cli, colunas="uf").status_code == 422
    assert consulta(cli, limite="muitos").status_code == 422


def test_filtros_numericos_e_booleanos_viram_texto(pasta):
    criar_base(pasta, "tipos", [{"ano": "2024", "ativo": "true"}, {"ano": "2023", "ativo": "false"}])
    assert local.consultar(pasta, "tipos", filtros={"ano": [2024]})["total"] == 1
    assert local.consultar(pasta, "tipos", filtros={"ativo": [True]})["total"] == 1
