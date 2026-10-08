"""Execução de consultas: validação de filtros, paginação automática e seleção de variáveis."""

from __future__ import annotations

import asyncio
import csv
import hashlib
import io
import json
import logging
from dataclasses import dataclass, field
from typing import Any

from .agregacao import Agregador
from .catalogo import Catalogo, Dataset, Parametro, montar_url
from .cliente import ClienteDataSUS
from .config import Config


log = logging.getLogger("api_sus")


class ErroConsulta(Exception):
    """Erro de validação da consulta feita pelo usuário (HTTP 422)."""


@dataclass
class Consulta:
    filtros: dict[str, Any] = field(default_factory=dict)
    colunas: list[str] | None = None
    filtros_locais: dict[str, Any] | None = None
    max_registros: int | None = None  # 0 = todos (até o teto configurado)
    paginar: bool = True
    agrupar_por: list[str] = field(default_factory=list)
    somar: list[str] = field(default_factory=list)


@dataclass
class Resultado:
    dataset: str
    url: str
    filtros: dict[str, Any]
    colunas: list[str]
    paginas_consultadas: int
    total: int
    avisos: list[str]
    registros: list[dict[str, Any]]
    registros_buscados: int = 0
    agregado: dict[str, Any] | None = None

    def como_dict(self) -> dict[str, Any]:
        d = {
            "dataset": self.dataset,
            "url_origem": self.url,
            "filtros": self.filtros,
            "colunas": self.colunas,
            "paginas_consultadas": self.paginas_consultadas,
            "registros_buscados": self.registros_buscados,
            "total": self.total,
            "avisos": self.avisos,
        }
        if self.agregado is not None:
            d["agregado"] = self.agregado
        d["dados"] = self.registros
        return d


# --------------------------------------------------------------------------- #
# Registros
# --------------------------------------------------------------------------- #


def extrair_registros(dados: Any, chave_lista: str | None = None) -> list[dict[str, Any]]:
    """Normaliza a resposta da API em uma lista de registros (dicts)."""
    if isinstance(dados, list):
        return [r if isinstance(r, dict) else {"valor": r} for r in dados]
    if isinstance(dados, dict):
        if chave_lista and isinstance(dados.get(chave_lista), list):
            return extrair_registros(dados[chave_lista])
        if len(dados) == 1:
            unico = next(iter(dados.values()))
            if isinstance(unico, list):
                return extrair_registros(unico)
        return [dados]
    if dados is None:
        return []
    return [{"valor": dados}]


def achatar(registro: dict[str, Any], prefixo: str = "") -> dict[str, Any]:
    """{"a": {"b": 1}} -> {"a.b": 1}. Listas viram texto JSON."""
    saida: dict[str, Any] = {}
    for chave, valor in registro.items():
        nome = f"{prefixo}{chave}"
        if isinstance(valor, dict):
            saida.update(achatar(valor, nome + "."))
        elif isinstance(valor, list):
            saida[nome] = json.dumps(valor, ensure_ascii=False)
        else:
            saida[nome] = valor
    return saida


def colunas_presentes(registros: list[dict[str, Any]]) -> list[str]:
    vistas: dict[str, None] = {}
    for r in registros:
        for k in achatar(r):
            vistas.setdefault(k, None)
    return list(vistas)


def _impressao_digital(registros: list[dict[str, Any]]) -> str:
    return hashlib.sha1(json.dumps(registros, sort_keys=True, default=str).encode()).hexdigest()


# --------------------------------------------------------------------------- #
# Validação
# --------------------------------------------------------------------------- #


def _converter(param: Parametro, valor: Any) -> Any:
    if isinstance(valor, list):
        return [_converter(param, v) for v in valor]
    texto = str(valor).strip()
    try:
        if param.tipo == "integer":
            convertido: Any = int(texto)
        elif param.tipo == "number":
            convertido = float(texto)
        elif param.tipo == "boolean":
            if texto.lower() in {"true", "1", "sim", "s"}:
                convertido = "true"
            elif texto.lower() in {"false", "0", "nao", "não", "n"}:
                convertido = "false"
            else:
                raise ValueError(texto)
        else:
            convertido = texto
    except ValueError as exc:
        raise ErroConsulta(f"Filtro '{param.nome}' deve ser do tipo {param.tipo}; recebido '{valor}'.") from exc

    if param.enum and convertido not in param.enum and str(convertido) not in [str(e) for e in param.enum]:
        raise ErroConsulta(f"Filtro '{param.nome}' aceita apenas {param.enum}; recebido '{valor}'.")
    if isinstance(convertido, (int, float)) and not isinstance(convertido, bool):
        if param.minimo is not None and convertido < param.minimo:
            raise ErroConsulta(f"Filtro '{param.nome}' deve ser >= {param.minimo}.")
        if param.maximo is not None and convertido > param.maximo:
            raise ErroConsulta(f"Filtro '{param.nome}' deve ser <= {param.maximo}.")
    return convertido


def validar_filtros(ds: Dataset, filtros: dict[str, Any]) -> tuple[dict[str, Any], dict[str, Any]]:
    """Valida os filtros contra a especificação. Retorna (query, path)."""
    conhecidos = {p.nome: p for p in ds.parametros}
    desconhecidos = [k for k in filtros if k not in conhecidos]
    if desconhecidos:
        raise ErroConsulta(
            f"Filtro(s) inexistente(s) para '{ds.id}': {desconhecidos}. "
            f"Disponíveis: {sorted(conhecidos) or 'nenhum'}."
        )

    query: dict[str, Any] = {}
    path: dict[str, Any] = {}
    for nome, valor in filtros.items():
        if valor is None or valor == "" or valor == []:
            continue
        p = conhecidos[nome]
        (path if p.local == "path" else query)[nome] = _converter(p, valor)

    controlados = {p.nome for p in (ds.param_limit, ds.param_offset) if p}
    faltando = [
        p.nome
        for p in ds.parametros
        if p.obrigatorio and p.nome not in controlados and p.nome not in query and p.nome not in path
    ]
    if faltando:
        raise ErroConsulta(f"Filtro(s) obrigatório(s) ausente(s) para '{ds.id}': {faltando}.")
    return query, path


# --------------------------------------------------------------------------- #
# Execução
# --------------------------------------------------------------------------- #


def _tamanho_pagina(config: Config, ds: Dataset, query: dict[str, Any]) -> int:
    lim = ds.param_limit
    assert lim is not None
    if lim.nome in query:
        return int(query[lim.nome])
    if lim.maximo:
        return int(lim.maximo)
    if lim.padrao:
        return int(lim.padrao)
    return config.tamanho_pagina_padrao


def _campos_para_api(ds: Dataset, colunas: list[str] | None, query: dict[str, Any]) -> None:
    """Se a base aceita o parâmetro "campos", pede à API só as colunas escolhidas."""
    param = ds.parametro("campos")
    if not colunas or param is None or param.local != "query" or "campos" in query:
        return
    nomes = list(dict.fromkeys(c.split(".")[0] for c in colunas))
    query["campos"] = ",".join(nomes)


async def executar(
    cliente: ClienteDataSUS, config: Config, catalogo: Catalogo, ds: Dataset, consulta: Consulta
) -> Resultado:
    query, path = validar_filtros(ds, consulta.filtros)
    agrupar_por, somar = list(consulta.agrupar_por or []), list(consulta.somar or [])
    if consulta.colunas:
        _campos_para_api(ds, [*consulta.colunas, *agrupar_por, *somar], query)
    url = montar_url(config, catalogo, ds, path)
    if consulta.max_registros is not None and consulta.max_registros < 0:
        raise ErroConsulta("max_registros deve ser >= 0 (0 = todos).")
    pedido = config.max_registros_teto if consulta.max_registros == 0 else consulta.max_registros
    max_registros = min(pedido or config.max_registros_padrao, config.max_registros_teto)

    avisos: list[str] = []
    registros: list[dict[str, Any]] = []
    agregador = Agregador(agrupar_por, somar) if (agrupar_por or somar) else None
    alvo = _alvo_local(consulta.filtros_locais) if consulta.filtros_locais else None
    vistas: dict[str, None] = {}
    contagem = {"buscados": 0, "total": 0}
    guardar = True

    def consumir(pagina: list[dict[str, Any]]) -> None:
        nonlocal guardar
        for r in pagina[: max_registros - contagem["buscados"]]:
            contagem["buscados"] += 1
            plano = achatar(r)
            if alvo and not _casa(plano, alvo):
                continue
            contagem["total"] += 1
            for k in plano:
                vistas.setdefault(k, None)
            if agregador:
                agregador.adicionar(plano)
            if guardar:
                if len(registros) >= config.max_linhas_detalhe:
                    guardar = False
                    registros.clear()
                    avisos.append(
                        f"Mais de {config.max_linhas_detalhe} linhas: os dados detalhados foram omitidos "
                        "(o resumo agregado considera todas)."
                    )
                else:
                    registros.append({c: plano.get(c) for c in consulta.colunas} if consulta.colunas else r)

    paginas = 0
    if not ds.paginavel or not consulta.paginar:
        pagina = extrair_registros(await cliente.get_json(url, query), ds.chave_lista)
        paginas = 1
        if len(pagina) > max_registros:
            avisos.append(f"Resultado truncado em {max_registros} registros.")
        consumir(pagina)
    else:
        lim, off = ds.param_limit, ds.param_offset
        assert lim is not None and off is not None
        tamanho = _tamanho_pagina(config, ds, query)
        # a maioria das bases começa em 0; algumas (ex.: "pagina") começam em 1 (default da spec)
        inicio = int(query.pop(off.nome, None) or off.padrao or 0)
        anterior = None
        while contagem["buscados"] < max_registros and paginas < config.max_paginas_teto:
            deslocamento = inicio + (paginas if config.modo_offset == "pagina" else paginas * tamanho)
            params = {**query, lim.nome: tamanho, off.nome: deslocamento}
            pagina = extrair_registros(await cliente.get_json(url, params), ds.chave_lista)
            paginas += 1
            if not pagina:
                break
            digital = _impressao_digital(pagina)
            if digital == anterior:
                avisos.append(
                    "A API devolveu a mesma página duas vezes; a paginação foi interrompida. "
                    "Verifique DATASUS_MODO_OFFSET ('pagina' ou 'registro')."
                )
                break
            anterior = digital
            consumir(pagina)
            if paginas % 20 == 0:
                log.info("%s: %d páginas, %d registros", ds.id, paginas, contagem["buscados"])
            if len(pagina) < tamanho:
                break
            if config.pausa_entre_paginas:
                await asyncio.sleep(config.pausa_entre_paginas)
        else:
            if contagem["buscados"] >= max_registros:
                avisos.append(
                    f"Limite de {max_registros} registros atingido; pode haver mais dados. "
                    "Aumente max_registros (0 = todos) para buscar mais."
                )
            else:
                avisos.append(f"Limite de {config.max_paginas_teto} páginas atingido; pode haver mais dados.")

    if consulta.colunas:
        ausentes = [c for c in consulta.colunas if c not in vistas]
        if ausentes and contagem["total"]:
            avisos.append(f"Coluna(s) não encontrada(s) nos dados: {ausentes}.")
        colunas = list(consulta.colunas)
    else:
        colunas = list(vistas)

    agregado = None
    if agregador:
        agregado = agregador.resultado()
        avisos.extend(agregado.pop("avisos"))

    return Resultado(
        dataset=ds.id,
        url=url,
        filtros={**path, **query},
        colunas=colunas,
        paginas_consultadas=paginas,
        total=contagem["total"],
        avisos=avisos,
        registros=registros,
        registros_buscados=contagem["buscados"],
        agregado=agregado,
    )


def _alvo_local(filtros: dict[str, Any]) -> dict[str, set[str]]:
    return {k: {_normal(x) for x in (v if isinstance(v, list) else [v])} for k, v in filtros.items()}


def _normal(v: Any) -> str:
    return str(v).strip().lower()


def _casa(plano: dict[str, Any], alvo: dict[str, set[str]]) -> bool:
    return all(_normal(plano.get(k)) in valores for k, valores in alvo.items())


async def amostrar_variaveis(
    cliente: ClienteDataSUS, config: Config, catalogo: Catalogo, ds: Dataset, filtros: dict[str, Any]
) -> dict[str, Any]:
    """Consulta poucos registros para descobrir as variáveis reais e exemplos de valores."""
    filtros = dict(filtros)
    if ds.param_limit and ds.param_limit.nome not in filtros:
        filtros[ds.param_limit.nome] = 5
    if ds.param_offset and ds.param_offset.nome not in filtros:
        filtros[ds.param_offset.nome] = ds.param_offset.padrao or 0
    res = await executar(cliente, config, catalogo, ds, Consulta(filtros=filtros, max_registros=5, paginar=False))
    exemplos: dict[str, Any] = {}
    for r in res.registros:
        for k, v in achatar(r).items():
            if k not in exemplos and v not in (None, ""):
                exemplos[k] = v
    return {"variaveis": res.colunas, "exemplos": exemplos, "registros_amostrados": res.total}


# --------------------------------------------------------------------------- #
# Saída
# --------------------------------------------------------------------------- #


def para_csv(resultado: Resultado, separador: str = ",") -> str:
    """CSV dos dados detalhados."""
    buffer = io.StringIO()
    escritor = csv.DictWriter(
        buffer, fieldnames=resultado.colunas, delimiter=separador, extrasaction="ignore", lineterminator="\n"
    )
    escritor.writeheader()
    for r in resultado.registros:
        escritor.writerow({k: ("" if v is None else v) for k, v in achatar(r).items()})
    return "﻿" + buffer.getvalue()  # BOM para o Excel reconhecer UTF-8


def para_xlsx(resultado: Resultado) -> bytes:
    """Gera uma planilha Excel com os dados e uma aba com informações da consulta."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font

    wb = Workbook(write_only=True)
    negrito = Font(bold=True)

    def cabecalho(aba, nomes):
        celulas = []
        for n in nomes:
            c = WriteOnlyCell(aba, value=n)
            c.font = negrito
            celulas.append(c)
        aba.append(celulas)

    if resultado.agregado is not None:
        ag = resultado.agregado
        resumo = wb.create_sheet("resumo")
        cabecalho(resumo, ag["colunas"])
        for linha in ag["linhas"]:
            resumo.append([linha.get(c) for c in ag["colunas"]])
        total = []
        for v in _linha_total(ag):
            c = WriteOnlyCell(resumo, value=v)
            c.font = negrito
            total.append(c)
        resumo.append(total)

    dados = wb.create_sheet("dados")
    cabecalho(dados, resultado.colunas)
    for r in resultado.registros:
        plano = achatar(r)
        dados.append([plano.get(c) for c in resultado.colunas])

    info = wb.create_sheet("consulta")
    linhas = [
        ("Base", resultado.dataset),
        ("Origem", resultado.url),
        ("Filtros", json.dumps(resultado.filtros, ensure_ascii=False)),
        ("Registros buscados na API", resultado.registros_buscados),
        ("Registros", resultado.total),
        ("Páginas consultadas", resultado.paginas_consultadas),
        *[("Aviso", a) for a in resultado.avisos],
    ]
    for rotulo, valor in linhas:
        c = WriteOnlyCell(info, value=rotulo)
        c.font = negrito
        info.append([c, valor])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def para_csv_agregado(resultado: Resultado, separador: str = ",") -> str:
    """CSV do resumo agregado (com linha de TOTAL)."""
    assert resultado.agregado is not None
    ag = resultado.agregado
    buffer = io.StringIO()
    escritor = csv.writer(buffer, delimiter=separador, lineterminator="\n")
    escritor.writerow(ag["colunas"])
    for linha in ag["linhas"]:
        escritor.writerow(["" if linha.get(c) is None else linha.get(c) for c in ag["colunas"]])
    escritor.writerow(_linha_total(ag))
    return "\ufeff" + buffer.getvalue()


def _linha_total(ag: dict[str, Any]) -> list[Any]:
    linha = ["TOTAL" if i == 0 else "" for i in range(len(ag["agrupar_por"]))]
    return linha + [ag["totais"][c] for c in ag["colunas"][len(ag["agrupar_por"]):]]
