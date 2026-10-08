"""Consultas sobre as bases baixadas para o computador (app local), com o DuckDB.

As bases ficam na pasta de dados do app:

    bases/<nome>.parquet   dados (todas as colunas como texto; colunas aninhadas achatadas, ex.: "endereco.uf")
    bases/<nome>.json      metadados (base de origem, filtros do download, registros, colunas...)
    exportacoes/           cópias das planilhas exportadas

Cada função abre uma conexão DuckDB em memória, lê o Parquet direto do disco (sem carregar a base
inteira na memória) e fecha a conexão no fim.

Segurança: o nome da base é conferido por RE_NOME_BASE (não há como apontar para fora da pasta);
os nomes de colunas vindos do usuário são conferidos contra o esquema real do Parquet e citados
como identificadores; os valores vão sempre como parâmetros do DuckDB, nunca no texto do SQL.
"""

from __future__ import annotations

import contextlib
import csv
import json
import logging
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterator

from .agregacao import _limpo
from .consulta import ErroConsulta, _linha_total

log = logging.getLogger("api_sus")

RE_NOME_BASE = re.compile(r"^[a-z0-9][a-z0-9_-]{0,79}$")
PARCIAL = ".parcial"
FORMATOS = {"xlsx", "csv"}

LIMITE_CATEGORIAS_MAX = 10_000
LIMITE_CONSULTA_MAX = 50_000
# Acima disso, o resumo agregado devolvido em JSON traz só os primeiros grupos (a exportação traz todos).
MAX_GRUPOS_RESPOSTA = 20_000
# Bases maiores que isso usam contagem aproximada de valores distintos (bem mais rápida).
DISTINTOS_EXATOS_ATE = 2_000_000
# Uma coluna é "numérica" se pelo menos esta fração dos valores preenchidos for número.
FRACAO_NUMERICA = 0.95
# Uma aba do Excel tem 1.048.576 linhas: uma vai para o cabeçalho.
LINHAS_MAX_XLSX = 1_048_575
TEXTO_MAX_XLSX = 32_767  # caracteres por célula
LOTE = 10_000  # linhas lidas do DuckDB por vez ao exportar

# Caracteres que o str.strip() do Python remove (os mesmos de str.isspace()).
_ESPACOS = (
    "\t\n\x0b\x0c\r\x1c\x1d\x1e\x1f \x85\xa0       "
    "         　"
)
# Mesmas regras de app.agregacao.numero (só com dígitos ASCII), em expressões regulares do DuckDB (RE2).
_RE_MILHAR_BR = r"-?\d{1,3}(\.\d{3})+(,\d+)?"
_RE_DECIMAL_VIRGULA = r"-?\d+,\d+"
_RE_FLOAT_PYTHON = r"[+-]?((\d(_?\d)*)?\.\d(_?\d)*|\d(_?\d)*\.?)([eE][+-]?\d(_?\d)*)?"


def _macros_sql() -> list[str]:
    """Macros SQL (texto fixo, sem nada vindo do usuário) para converter texto em número como numero()."""
    espacos = " || ".join(f"chr({ord(c)})" for c in _ESPACOS)
    return [
        f"CREATE TEMP MACRO api_sus_aparar(v) AS trim(CAST(v AS VARCHAR), {espacos})",
        f"""CREATE TEMP MACRO api_sus_numero_texto(t) AS (
            CASE
                WHEN regexp_full_match(t, '{_RE_MILHAR_BR}')
                    THEN TRY_CAST(replace(replace(t, '.', ''), ',', '.') AS DOUBLE)
                WHEN regexp_full_match(t, '{_RE_DECIMAL_VIRGULA}') THEN TRY_CAST(replace(t, ',', '.') AS DOUBLE)
                WHEN regexp_full_match(t, '{_RE_FLOAT_PYTHON}') THEN TRY_CAST(replace(t, '_', '') AS DOUBLE)
            END
        )""",
        "CREATE TEMP MACRO api_sus_finito(x) AS (CASE WHEN isfinite(x) THEN x END)",
        "CREATE TEMP MACRO api_sus_numero(v) AS api_sus_finito(api_sus_numero_texto(api_sus_aparar(v)))",
    ]


class NaoEncontrado(LookupError):
    """Base ou coluna inexistente (HTTP 404)."""


class BaseNaoEncontrada(NaoEncontrado):
    pass


class ColunaNaoEncontrada(NaoEncontrado):
    pass


# --------------------------------------------------------------------------- #
# Pastas e arquivos
# --------------------------------------------------------------------------- #


def pasta_bases(pasta: Path) -> Path:
    return Path(pasta) / "bases"


def pasta_exportacoes(pasta: Path) -> Path:
    return Path(pasta) / "exportacoes"


def validar_nome(nome: Any) -> str:
    """O nome da base, se for um nome válido (só [a-z0-9_-]; nunca um caminho)."""
    if not isinstance(nome, str) or not RE_NOME_BASE.fullmatch(nome):
        raise BaseNaoEncontrada(f"Base '{nome}' não encontrada. Consulte GET /api/local/bases.")
    return nome


def _parquet(pasta: Path, nome: str) -> Path:
    caminho = pasta_bases(pasta) / f"{validar_nome(nome)}.parquet"
    if not caminho.is_file():
        raise BaseNaoEncontrada(f"Base '{nome}' não encontrada. Consulte GET /api/local/bases.")
    return caminho


def _apagar(*caminhos: Path) -> None:
    for caminho in caminhos:
        with contextlib.suppress(OSError):
            caminho.unlink(missing_ok=True)


def citar(nome: str) -> str:
    """Nome de coluna como identificador SQL: "nome", com aspas internas duplicadas."""
    return '"' + nome.replace('"', '""') + '"'


def _duckdb():
    try:
        import duckdb
    except ImportError as exc:
        raise ErroConsulta("O pacote 'duckdb' não está instalado (pip install duckdb).") from exc
    return duckdb


def _conectar(duckdb) -> Any:
    """Conexão em memória, sem a barra de progresso que o DuckDB desenha no terminal em consultas longas."""
    con = duckdb.connect(":memory:")
    con.execute("SET enable_progress_bar = false")
    return con


# --------------------------------------------------------------------------- #
# Validação da entrada
# --------------------------------------------------------------------------- #


def _unicos(itens: list[str] | None, campo: str) -> list[str]:
    if itens is None:
        return []
    if isinstance(itens, str) or not isinstance(itens, (list, tuple)):
        raise ErroConsulta(f"'{campo}' deve ser uma lista de nomes de colunas.")
    vistos: dict[str, None] = {}
    for item in itens:
        if not isinstance(item, str):
            raise ErroConsulta(f"'{campo}' deve ser uma lista de nomes de colunas; recebido {item!r}.")
        vistos.setdefault(item, None)
    return list(vistos)


def _texto_filtro(coluna: str, valor: Any) -> str | None:
    """Valor de filtro como texto (as colunas da base local são texto); None = nulo."""
    if valor is None or isinstance(valor, str):
        return valor
    if isinstance(valor, bool):
        return "true" if valor else "false"
    if isinstance(valor, (int, float)):
        return str(valor)
    raise ErroConsulta(f"Valor inválido no filtro '{coluna}': {valor!r}. Use texto, número ou null.")


def normalizar_filtros(filtros: Any) -> dict[str, list[str | None]]:
    """{coluna: [valores]} (aceita também um valor solto no lugar da lista)."""
    if filtros is None:
        return {}
    if not isinstance(filtros, dict):
        raise ErroConsulta('Os filtros devem ser um objeto {"coluna": [valores]}.')
    saida: dict[str, list[str | None]] = {}
    for coluna, valores in filtros.items():
        if not isinstance(coluna, str):
            raise ErroConsulta(f"Nome de coluna inválido nos filtros: {coluna!r}.")
        lista = list(valores) if isinstance(valores, (list, tuple)) else [valores]
        unicos: dict[str | None, None] = {}
        for v in lista:
            unicos.setdefault(_texto_filtro(coluna, v), None)
        saida[coluna] = list(unicos)
    return saida


def _limite(valor: Any, maximo: int, minimo: int = 0, campo: str = "limite") -> int:
    if isinstance(valor, bool) or not isinstance(valor, int):
        raise ErroConsulta(f"'{campo}' deve ser um número inteiro.")
    if valor < minimo or valor > maximo:
        raise ErroConsulta(f"'{campo}' deve estar entre {minimo} e {maximo}; recebido {valor}.")
    return valor


# --------------------------------------------------------------------------- #
# Conexão
# --------------------------------------------------------------------------- #


@dataclass
class _Base:
    """Uma base aberta: a view "base" (todas as colunas como texto) e as macros de número."""

    con: Any
    nome: str
    colunas: list[str]

    def __post_init__(self) -> None:
        self._nomes = set(self.colunas)

    def coluna(self, nome: Any) -> str:
        """Identificador SQL da coluna, se ela existir no esquema da base."""
        if not isinstance(nome, str) or nome not in self._nomes:
            raise ColunaNaoEncontrada(f"Coluna '{nome}' não existe na base '{self.nome}'.")
        return citar(nome)

    def selecao(self, colunas: list[str] | None) -> list[str]:
        escolhidas = _unicos(colunas, "colunas")
        for c in escolhidas:
            self.coluna(c)
        return escolhidas or list(self.colunas)

    def onde(self, filtros: Any, ignorar: str | None = None) -> tuple[list[str], list[Any]]:
        """Condições (AND) e parâmetros dos filtros {coluna: [valores]}; "ignorar" pula uma coluna."""
        condicoes: list[str] = []
        parametros: list[Any] = []
        for coluna, valores in normalizar_filtros(filtros).items():
            ident = self.coluna(coluna)
            if coluna == ignorar or not valores:
                continue
            textos = [v for v in valores if v is not None]
            termos = []
            if textos:
                termos.append(f"{ident} IN ({', '.join('?' for _ in textos)})")
                parametros.extend(textos)
            if len(textos) < len(valores):
                termos.append(f"{ident} IS NULL")
            condicoes.append("(" + " OR ".join(termos) + ")")
        return condicoes, parametros

    def contar(self, condicoes: list[str] | None = None, parametros: list[Any] | None = None) -> int:
        return int(self.con.execute(f"SELECT count(*) FROM base{_where(condicoes)}", parametros or []).fetchone()[0])


def _where(condicoes: list[str] | None) -> str:
    return (" WHERE " + " AND ".join(condicoes)) if condicoes else ""


@contextlib.contextmanager
def _abrir(pasta: Path, nome: str) -> Iterator[_Base]:
    caminho = _parquet(pasta, nome)
    duckdb = _duckdb()
    con = _conectar(duckdb)
    try:
        try:
            relacao = con.read_parquet(str(caminho))  # caminho passado à API, não ao texto do SQL
            esquema = list(zip(relacao.columns, (str(t) for t in relacao.types)))
        except duckdb.Error as exc:
            raise ErroConsulta(f"Não foi possível ler a base '{nome}': {exc}") from exc
        # As bases gravadas pelo app já são só texto; o CAST protege contra arquivos de outra origem.
        projecao = ", ".join(
            citar(c) if tipo == "VARCHAR" else f"CAST({citar(c)} AS VARCHAR) AS {citar(c)}" for c, tipo in esquema
        )
        relacao.project(projecao).create_view("base")
        for macro in _macros_sql():
            con.execute(macro)
        yield _Base(con, nome, [c for c, _ in esquema])
    finally:
        con.close()


# --------------------------------------------------------------------------- #
# Bases
# --------------------------------------------------------------------------- #


def _ler_metadados(caminho: Path) -> tuple[dict[str, Any], str | None]:
    try:
        meta = json.loads(caminho.read_text(encoding="utf-8"))
    except (OSError, ValueError) as exc:
        return {}, f"Metadados ilegíveis ({exc.__class__.__name__})."
    return (meta, None) if isinstance(meta, dict) else ({}, "Metadados ilegíveis.")


def _contar_parquet(caminho: Path) -> int | None:
    duckdb = _duckdb()
    con = _conectar(duckdb)
    try:
        return int(con.execute("SELECT count(*) FROM read_parquet(?)", [str(caminho)]).fetchone()[0])
    except duckdb.Error:
        return None
    finally:
        con.close()


def _metadados(pasta: Path, nome: str) -> dict[str, Any]:
    """Metadados da base (com valores padrão para os campos que faltarem) e o tamanho em disco."""
    pb = pasta_bases(pasta)
    parquet, arquivo_json = pb / f"{nome}.parquet", pb / f"{nome}.json"
    meta, problema = _ler_metadados(arquivo_json)
    saida: dict[str, Any] = {
        "nome": nome,
        "dataset": None,
        "titulo": nome,
        "filtros": {},
        "registros": None,
        "colunas": [],
        "baixado_em": None,
        "url_origem": None,
        "avisos": [],
        "paginas": 0,
        **meta,
    }
    saida["nome"] = nome  # vale o nome do arquivo
    if problema:
        saida["avisos"] = [*saida["avisos"], problema]
    if not isinstance(saida["registros"], int):
        saida["registros"] = _contar_parquet(parquet)
    if not saida["baixado_em"]:
        mtime = parquet.stat().st_mtime
        saida["baixado_em"] = datetime.fromtimestamp(mtime, timezone.utc).isoformat(timespec="seconds")
    saida["tamanho_bytes"] = parquet.stat().st_size + arquivo_json.stat().st_size
    return saida


def _momento(meta: dict[str, Any]) -> float:
    try:
        quando = datetime.fromisoformat(str(meta.get("baixado_em")))
    except ValueError:
        return 0.0
    if quando.tzinfo is None:
        quando = quando.replace(tzinfo=timezone.utc)
    return quando.timestamp()


def listar_bases(pasta: Path) -> list[dict[str, Any]]:
    """Bases baixadas (metadados + "tamanho_bytes"), mais recentes primeiro."""
    pb = pasta_bases(pasta)
    pb.mkdir(parents=True, exist_ok=True)
    pasta_exportacoes(pasta).mkdir(parents=True, exist_ok=True)
    bases = []
    for arquivo_json in pb.glob("*.json"):
        nome = arquivo_json.name[: -len(".json")]
        if not RE_NOME_BASE.fullmatch(nome) or not (pb / f"{nome}.parquet").is_file():
            continue  # arquivos parciais, de outros programas ou base sem dados
        try:
            bases.append(_metadados(pasta, nome))
        except OSError as exc:  # apagada enquanto era listada
            log.warning("Base local '%s' ignorada: %s", nome, exc)
    bases.sort(key=lambda m: (_momento(m), m["nome"]), reverse=True)
    return bases


def excluir_base(pasta: Path, nome: str) -> str:
    """Apaga os arquivos da base (dados e metadados)."""
    pb = pasta_bases(pasta)
    validar_nome(nome)
    arquivos = [pb / f"{nome}.parquet", pb / f"{nome}.json"]
    if not any(a.is_file() for a in arquivos):
        raise BaseNaoEncontrada(f"Base '{nome}' não encontrada. Consulte GET /api/local/bases.")
    for arquivo in arquivos:
        try:
            arquivo.unlink(missing_ok=True)
        except OSError as exc:
            raise ErroConsulta(
                f"Não foi possível apagar '{arquivo.name}' ({exc.strerror or exc}). "
                "Feche os programas que estejam usando o arquivo e tente de novo."
            ) from exc
    return nome


# --------------------------------------------------------------------------- #
# Colunas e categorias
# --------------------------------------------------------------------------- #


def colunas(pasta: Path, nome: str) -> dict[str, Any]:
    """Para cada coluna: valores distintos, nulos, se é numérica e um exemplo de valor."""
    with _abrir(pasta, nome) as base:
        registros = base.contar()
        aproximado = registros > DISTINTOS_EXATOS_ATE
        partes = []
        for c in base.colunas:
            ident = citar(c)
            distintos = f"approx_count_distinct({ident})" if aproximado else f"count(DISTINCT {ident})"
            partes += [
                distintos,
                f"count({ident})",
                f"count(*) FILTER (WHERE api_sus_aparar({ident}) <> '')",
                f"count(api_sus_numero({ident}))",
                f"coalesce(any_value({ident}) FILTER (WHERE api_sus_aparar({ident}) <> ''), any_value({ident}))",
            ]
        linha = base.con.execute(f"SELECT {', '.join(partes)} FROM base").fetchone() if partes else ()

    saida = []
    for i, c in enumerate(base.colunas):
        distintos, nao_nulos, preenchidos, numericos, exemplo = linha[i * 5:(i + 1) * 5]
        saida.append(
            {
                "nome": c,
                "distintos": int(distintos),
                "nulos": registros - int(nao_nulos),
                "numerica": preenchidos > 0 and numericos >= FRACAO_NUMERICA * preenchidos,
                "exemplo": exemplo,
            }
        )
    return {"nome": nome, "registros": registros, "distintos_aproximados": aproximado, "colunas": saida}


def categorias(
    pasta: Path,
    nome: str,
    coluna: str,
    busca: str | None = None,
    limite: int = 200,
    filtros: Any = None,
) -> dict[str, Any]:
    """Valores da coluna e quantos registros têm cada um (mais frequentes primeiro).

    Aplica os filtros das OUTRAS colunas e ignora o da própria coluna (filtros em cascata: a lista
    mostra o que ainda pode ser escolhido). ``busca`` procura um trecho, sem diferenciar maiúsculas
    nem acentos.
    """
    limite = _limite(limite, LIMITE_CATEGORIAS_MAX, minimo=1)
    with _abrir(pasta, nome) as base:
        ident = base.coluna(coluna)
        condicoes, parametros = base.onde(filtros, ignorar=coluna)
        termo = (busca or "").strip()
        if termo:
            condicoes.append(f"contains(strip_accents(lower({ident})), strip_accents(lower(?)))")
            parametros.append(termo)
        sql = (
            "SELECT valor, registros, count(*) OVER () AS total FROM ("
            f"SELECT {ident} AS valor, count(*) AS registros FROM base{_where(condicoes)} GROUP BY 1"
            ") ORDER BY registros DESC, valor ASC NULLS LAST LIMIT ?"
        )
        linhas = base.con.execute(sql, [*parametros, limite]).fetchall()

    total = int(linhas[0][2]) if linhas else 0
    return {
        "coluna": coluna,
        "total_distintos": total,
        "valores": [{"valor": v, "registros": int(n)} for v, n, _ in linhas],
        "truncado": total > len(linhas),
    }


# --------------------------------------------------------------------------- #
# Consulta e agregação
# --------------------------------------------------------------------------- #


def _agregar(
    base: _Base,
    condicoes: list[str],
    parametros: list[Any],
    agrupar_por: list[str],
    somar: list[str],
    max_linhas: int | None = None,
) -> dict[str, Any]:
    """Registros e somas por grupo, no formato de app.agregacao.Agregador.resultado() (sem "avisos")."""
    nomes_saida = [*agrupar_por, "registros", *(f"soma_{c}" for c in somar)]
    repetidos = sorted({n for n in nomes_saida if nomes_saida.count(n) > 1})
    if repetidos:
        raise ErroConsulta(f"Conflito de nomes no resumo: {repetidos}. Agrupe por outras colunas.")

    grupos_id = [base.coluna(c) for c in agrupar_por]
    somas_id = [base.coluna(c) for c in somar]
    selecao = [f"{ident} AS g{i}" for i, ident in enumerate(grupos_id)]
    selecao.append("count(*) AS registros")
    selecao += [f"coalesce(sum(api_sus_numero({ident})), 0) AS s{j}" for j, ident in enumerate(somas_id)]
    where = _where(condicoes)
    k = len(agrupar_por)

    def contagens(linha: tuple) -> dict[str, Any]:
        """{"registros", "soma_<col>"...} a partir de uma linha (grupos..., registros, somas...)."""
        saida: dict[str, Any] = {"registros": int(linha[k])}
        for j, c in enumerate(somar):
            saida[f"soma_{c}"] = _limpo(linha[k + 1 + j])
        return saida

    if not agrupar_por:  # um único grupo com todos os registros
        totais = contagens(base.con.execute(f"SELECT {', '.join(selecao)} FROM base{where}", parametros).fetchone())
        linhas = [dict(totais)] if totais["registros"] else []
        return _formato_agregado(agrupar_por, somar, nomes_saida, len(linhas), linhas, totais)

    lista_grupos = ", ".join(grupos_id)
    ordem = ["nivel DESC"]
    ordem += ["s0 DESC"] if somar else []
    ordem += ["registros DESC", *(f"g{i} ASC NULLS LAST" for i in range(len(agrupar_por)))]
    sql = (
        f"WITH r AS (SELECT {', '.join(selecao)}, GROUPING({lista_grupos}) AS nivel FROM base{where} "
        f"GROUP BY GROUPING SETS (({lista_grupos}), ())) "
        f"SELECT *, count(*) FILTER (WHERE nivel = 0) OVER () AS n_grupos FROM r ORDER BY {', '.join(ordem)}"
    )
    argumentos = list(parametros)
    if max_linhas is not None:
        sql += " LIMIT ?"
        argumentos.append(max_linhas + 1)  # + a linha de totais
    resultado = base.con.execute(sql, argumentos).fetchall()

    total_linha, *grupos = resultado  # a 1ª linha é a dos totais (GROUPING SETS "()")
    linhas = [{**dict(zip(agrupar_por, linha[:k])), **contagens(linha)} for linha in grupos]
    n_grupos = int(total_linha[-1])
    return _formato_agregado(agrupar_por, somar, nomes_saida, n_grupos, linhas, contagens(total_linha))


def _formato_agregado(agrupar_por, somar, colunas_saida, grupos, linhas, totais) -> dict[str, Any]:
    return {
        "agrupar_por": list(agrupar_por),
        "somar": list(somar),
        "colunas": list(colunas_saida),
        "grupos": grupos,
        "linhas": linhas,
        "totais": totais,
    }


@dataclass
class _Pedido:
    """Consulta já validada contra o esquema da base."""

    condicoes: list[str]
    parametros: list[Any]
    colunas: list[str]
    agrupar_por: list[str]
    somar: list[str]
    filtros: dict[str, list[str | None]]

    @property
    def agrega(self) -> bool:
        return bool(self.agrupar_por or self.somar)


def _pedido(base: _Base, filtros: Any, colunas: Any, agrupar_por: Any, somar: Any) -> _Pedido:
    condicoes, parametros = base.onde(filtros)
    agrupar = _unicos(agrupar_por, "agrupar_por")
    soma = _unicos(somar, "somar")
    for c in [*agrupar, *soma]:
        base.coluna(c)
    return _Pedido(condicoes, parametros, base.selecao(colunas), agrupar, soma, normalizar_filtros(filtros))


def consultar(
    pasta: Path,
    nome: str,
    filtros: Any = None,
    colunas: list[str] | None = None,
    agrupar_por: list[str] | None = None,
    somar: list[str] | None = None,
    limite: int = 500,
) -> dict[str, Any]:
    """Registros filtrados (até ``limite``), o total e, se pedido, o resumo agregado."""
    limite = _limite(limite, LIMITE_CONSULTA_MAX)
    with _abrir(pasta, nome) as base:
        p = _pedido(base, filtros, colunas, agrupar_por, somar)
        total = base.contar(p.condicoes, p.parametros)
        linhas = []
        if limite:
            sql = f"SELECT {', '.join(citar(c) for c in p.colunas)} FROM base{_where(p.condicoes)} LIMIT ?"
            linhas = [dict(zip(p.colunas, r)) for r in base.con.execute(sql, [*p.parametros, limite]).fetchall()]
        agregado = (
            _agregar(base, p.condicoes, p.parametros, p.agrupar_por, p.somar, MAX_GRUPOS_RESPOSTA) if p.agrega else None
        )
    return {"total": total, "colunas": p.colunas, "linhas": linhas, "agregado": agregado}


# --------------------------------------------------------------------------- #
# Exportação
# --------------------------------------------------------------------------- #


def _lotes(base: _Base, sql: str, parametros: list[Any]) -> Iterator[list[tuple]]:
    cursor = base.con.execute(sql, parametros)
    while True:
        lote = cursor.fetchmany(LOTE)
        if not lote:
            return
        yield lote


def _reservar_arquivo(pasta: Path, nome: str, formato: str) -> Path:
    """exportacoes/<nome>__<AAAAMMDD-HHMMSS>.<ext> (com "-2", "-3"... se já existir)."""
    destino = pasta_exportacoes(pasta)
    destino.mkdir(parents=True, exist_ok=True)
    raiz = f"{nome}__{datetime.now().strftime('%Y%m%d-%H%M%S')}"
    n = 1
    while True:
        arquivo = destino / f"{raiz}{'' if n == 1 else f'-{n}'}.{formato}"
        if not arquivo.exists():
            try:  # cria o parcial de forma exclusiva: duas exportações no mesmo segundo não colidem
                Path(f"{arquivo}{PARCIAL}").open("x").close()
                return arquivo
            except FileExistsError:
                pass
        n += 1


def exportar(
    pasta: Path,
    nome: str,
    formato: str,
    filtros: Any = None,
    colunas: list[str] | None = None,
    agrupar_por: list[str] | None = None,
    somar: list[str] | None = None,
) -> tuple[Path, str]:
    """Gera a planilha (xlsx ou csv) com TODOS os registros filtrados e guarda uma cópia em exportacoes/.

    Retorna (caminho do arquivo, nome sugerido para o download).
    """
    formato = str(formato or "").strip().lower()
    if formato not in FORMATOS:
        raise ErroConsulta(f"Formato '{formato}' inválido. Use 'xlsx' ou 'csv'.")
    with _abrir(pasta, nome) as base:
        p = _pedido(base, filtros, colunas, agrupar_por, somar)
        agregado = _agregar(base, p.condicoes, p.parametros, p.agrupar_por, p.somar) if p.agrega else None
        arquivo = _reservar_arquivo(pasta, nome, formato)
        parcial = Path(f"{arquivo}{PARCIAL}")
        try:
            if formato == "csv":
                _escrever_csv(parcial, base, p, agregado)
            else:
                info = _InfoExportacao(
                    meta=_ler_metadados(pasta_bases(pasta) / f"{nome}.json")[0],
                    registros_base=base.contar(),
                    total=base.contar(p.condicoes, p.parametros),
                )
                _escrever_xlsx(parcial, base, p, agregado, info)
            parcial.replace(arquivo)
        except BaseException:
            _apagar(parcial)
            raise
    log.info("Base local '%s' exportada para %s", nome, arquivo)
    return arquivo, arquivo.name


def _escrever_csv(caminho: Path, base: _Base, p: _Pedido, agregado: dict[str, Any] | None) -> None:
    """Com resumo agregado, exporta o resumo (com a linha TOTAL); senão, os dados. ";" e UTF-8 com BOM."""
    with caminho.open("w", encoding="utf-8-sig", newline="") as arquivo:
        escritor = csv.writer(arquivo, delimiter=";", lineterminator="\n")
        if agregado is not None:
            escritor.writerow(agregado["colunas"])
            for linha in agregado["linhas"]:
                escritor.writerow(["" if linha.get(c) is None else linha.get(c) for c in agregado["colunas"]])
            escritor.writerow(_linha_total(agregado))
            return
        escritor.writerow(p.colunas)
        sql = f"SELECT {', '.join(citar(c) for c in p.colunas)} FROM base{_where(p.condicoes)}"
        for lote in _lotes(base, sql, p.parametros):
            escritor.writerows(["" if v is None else v for v in r] for r in lote)


def _json(valor: Any) -> str:
    return json.dumps(valor, ensure_ascii=False)


@dataclass
class _InfoExportacao:
    meta: dict[str, Any]
    registros_base: int
    total: int


class _CelulasXlsx:
    """Converte valores para células seguras: sem caracteres proibidos, sem virar fórmula."""

    def __init__(self):
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
        from openpyxl.styles import Font

        self._celula = WriteOnlyCell
        self._proibidos = ILLEGAL_CHARACTERS_RE
        self.negrito = Font(bold=True)
        self.cortados = 0

    def _texto(self, valor: str) -> str:
        if self._proibidos.search(valor):
            valor = self._proibidos.sub("", valor)
        if len(valor) > TEXTO_MAX_XLSX:
            valor = valor[:TEXTO_MAX_XLSX]
            self.cortados += 1
        return valor

    def _nova(self, aba, valor: Any):
        celula = self._celula(aba, value=valor)
        if isinstance(valor, str) and valor.startswith("="):
            celula.data_type = "s"  # texto, não fórmula
        return celula

    def valor(self, aba, valor: Any) -> Any:
        if not isinstance(valor, str):
            return valor
        valor = self._texto(valor)
        return self._nova(aba, valor) if valor.startswith("=") else valor

    def linha(self, aba, valores) -> list[Any]:
        return [self.valor(aba, v) for v in valores]

    def destaque(self, aba, valores) -> list[Any]:
        """Células em negrito (cabeçalhos, linha TOTAL e rótulos)."""
        celulas = []
        for v in valores:
            celula = self._nova(aba, self._texto(v) if isinstance(v, str) else v)
            celula.font = self.negrito
            celulas.append(celula)
        return celulas


def _escrever_xlsx(
    caminho: Path, base: _Base, p: _Pedido, agregado: dict[str, Any] | None, info: _InfoExportacao
) -> None:
    """Abas "resumo" (se houver agregado, com linha TOTAL em negrito), "dados" e "consulta"."""
    from openpyxl import Workbook

    wb = Workbook(write_only=True)
    celulas = _CelulasXlsx()
    avisos: list[str] = []

    if agregado is not None:
        resumo = wb.create_sheet("resumo")
        resumo.append(celulas.destaque(resumo, agregado["colunas"]))
        linhas_resumo = agregado["linhas"]
        if len(linhas_resumo) > LINHAS_MAX_XLSX - 1:  # - a linha TOTAL
            linhas_resumo = linhas_resumo[: LINHAS_MAX_XLSX - 1]
            avisos.append(
                f"O resumo tem {agregado['grupos']} grupos; a aba 'resumo' mostra só os "
                f"{LINHAS_MAX_XLSX - 1} primeiros (limite do Excel). O TOTAL considera todos."
            )
        for linha in linhas_resumo:
            resumo.append(celulas.linha(resumo, [linha.get(c) for c in agregado["colunas"]]))
        resumo.append(celulas.destaque(resumo, _linha_total(agregado)))

    dados = wb.create_sheet("dados")
    dados.append(celulas.destaque(dados, p.colunas))
    sql = f"SELECT {', '.join(citar(c) for c in p.colunas)} FROM base{_where(p.condicoes)} LIMIT ?"
    exportadas = 0
    for lote in _lotes(base, sql, [*p.parametros, LINHAS_MAX_XLSX]):
        for r in lote:
            dados.append(celulas.linha(dados, r))
        exportadas += len(lote)
    if info.total > exportadas:
        avisos.append(
            f"A consulta tem {info.total} registros, mas uma aba do Excel comporta {LINHAS_MAX_XLSX}: a aba "
            f"'dados' foi cortada em {exportadas} linhas. Exporte em CSV para obter todos os registros."
        )
    if celulas.cortados:
        avisos.append(
            f"{celulas.cortados} valor(es) com mais de {TEXTO_MAX_XLSX} caracteres foram cortados (limite do Excel)."
        )

    aba = wb.create_sheet("consulta")
    meta = info.meta
    linhas_info: list[tuple[str, Any]] = [
        ("Base", base.nome),
        ("Título", meta.get("titulo")),
        ("Base de origem", meta.get("dataset")),
        ("Filtros do download", _json(meta.get("filtros") or {})),
        ("Baixada em", meta.get("baixado_em")),
        ("Origem", meta.get("url_origem")),
        ("Filtros", _json(p.filtros)),
        ("Colunas", ", ".join(p.colunas)),
        ("Agrupar por", ", ".join(p.agrupar_por)),
        ("Somar", ", ".join(p.somar)),
        ("Exportada em", datetime.now().astimezone().isoformat(timespec="seconds")),
        ("Registros na base", info.registros_base),
        ("Registros filtrados", info.total),
        ("Linhas na aba 'dados'", exportadas),
    ]
    if agregado is not None:
        linhas_info.append(("Grupos", agregado["grupos"]))
    linhas_info += [("Aviso", a) for a in avisos]
    for rotulo, valor in linhas_info:
        aba.append([*celulas.destaque(aba, [rotulo]), celulas.valor(aba, valor)])

    wb.save(str(caminho))
